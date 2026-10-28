"""Write the reviewer spreadsheet for Hindi and Kannada readers: one row per phrase, a verdict dropdown, a correction column.

  PYTHONPATH=<site-packages with openpyxl> python scripts/build_review_sheet.py     # needs NEBIUS_API_KEY for the AI-text sheets (they are skipped without it)

Output: my_data/hindi_kannada_review.xlsx (git-ignored: it is for sending to reviewers, not for the repository). The markdown pack (docs/language_review.md)
holds the same material for people who read the repository. Sheets: Start here, Hindi fixed text, Hindi AI text, Kannada fixed text,
Kannada medicine table, Medicine map (a medical check, in English), Kannada AI text (switched off).
"""
import datetime, pathlib, sys
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from dawa_buddy import i18n, llm, medinfo
from dawa_buddy.explainer import PHRASES, dose_text, explain
from dawa_buddy.reader import read_offline
from dawa_buddy.scheduler import build_schedule

RX = read_offline("""Tab Telma 40 mg 1-0-0 x 30 days
Tab Dolo 650 mg 1-0-1 x 5 days after food
Tab Atorva 10 mg 0-0-1 x 30 days
Tab Metformin 500 mg 1-1-1 x 30 days after food
Tab Pantop 40 mg 1-0-0 x 30 days before food
Tab Cetzine 10 mg SOS
Tab Ecosprin 75 mg 0-0-1 to continue""")
TEAL, TINT = "0F6B5C", "EAF4F0"
VERDICTS = '"OK,Change it,Not sure"'
YN = '"Yes,No,Not sure"'


def flat(d, prefix=""):
    for k, v in d.items():
        if isinstance(v, dict) and not (set(v) >= {"en", "hi", "kn"} or set(v) >= {"en", "kn"}):
            yield from flat(v, f"{prefix}{k}.")
        else:
            yield f"{prefix}{k}", v


NICE = {  # plain descriptions for the technical keys, so a reviewer never sees "slots.m"
    "slots.m": "Dose sentence: the word 'morning'", "slots.a": "Dose sentence: the word 'afternoon'", "slots.n": "Dose sentence: the word 'night'",
    "unit": "Dose sentence: 'tablet' (one)", "units": "Dose sentence: 'tablets' (more than one)", "dose": "Dose sentence: 'dose' (for a syrup or injection)", "doses": "Dose sentence: 'doses'",
    "times.1": "Dose sentence start: once a day", "times.2": "Dose sentence start: twice a day", "times.3": "Dose sentence start: 3 times a day",
    "before": "Dose sentence: before food", "after": "Dose sentence: after food", "bedtime": "Dose sentence: at bedtime", "days": "Dose sentence: 'for N days' (N is filled in)",
    "end": "Dose sentence: how a sentence ends", "sos": "Sentence for a medicine taken only when needed", "go_on": "Sentence added when the prescription says to continue",
    "days_missing": "Sentence added when the number of days is not written", "unclear": "Sentence when the dose could not be read", "stat": "Sentence for a single, one-time dose",
    "slots.morning": "Timetable: 'Morning'", "slots.afternoon": "Timetable: 'Afternoon'", "slots.night": "Timetable: 'Night'", "slots.as_needed": "Timetable: 'As needed'",
    "before": "Dose sentence and timetable: before food", "as_written": "Timetable: shown when nothing more is written", "sos": "Sentence for a medicine taken only when needed",
    "share_title": "Title of the timetable people share", "share_days": "Shared timetable: duration line ({n} is filled in)", "share_foot": "Shared timetable: closing line"}


def fixed_rows(lang):
    """(part, where you see it, English, text in `lang`) for every fixed phrase a reader of `lang` can meet."""
    rows = []
    en_p, p = PHRASES["en"], PHRASES[lang]
    for key, v in flat(p):
        e = dict(flat(en_p)).get(key)
        rows.append(("Dose sentence piece", NICE.get(key, f"Dose sentence: {key}"), e, v))
    for ex_ in RX.medications:
        rows.append(("Whole dose sentence", f"{ex_.name}: prescription says {ex_.frequency or '?'}", dose_text(ex_, "en"), dose_text(ex_, lang)))
    where = {"SEVERITY": "Heading above a group of safety notes", "TITLES": "Title of a safety note", "DETAILS": "Sentence of a safety note ({...} is filled with a name or number)",
             "NOT_WRITTEN": "Part of the 'something is not written' note", "STATUS": "Regulator's batch-alert wording", "LABELS": "Heading or caption on the explainer screen",
             "DISCLAIMER": "Under every explanation", "AI_NOTE": "Notice above the medicine cards (AI-written text)", "FIXED_NOTE": "Notice above the medicine cards (fixed list)",
             "PARTIAL_NOTE": "Notice above the medicine cards (some fixed, some English)", "FALLBACK_NOTE": "Notice above the medicine cards (English fallback)"}
    for name, w in where.items():
        tbl = getattr(i18n, name)
        items = tbl.items() if not (set(tbl) >= {"en", "hi", "kn"}) else [("", tbl)]
        for key, d in items:
            if d.get("en"):
                rows.append((name.replace("_", " ").title(), f"{w}" + (f": {key}" if key else ""), d["en"], d[lang]))
    for key, v in flat({k: x for k, x in i18n.TIMETABLE.items() if k != "headers"}):
        rows.append(("Timetable", NICE.get(key, f"Timetable: {key}"), v["en"], v[lang]))
    h = i18n.TIMETABLE["headers"]
    rows.append(("Timetable", "Timetable: column headings", " | ".join(h["en"]), " | ".join(h[lang])))
    sched_en, sched = build_schedule(RX, "en"), build_schedule(RX, lang)
    for a, b in zip(sched_en, sched):
        rows.append(("Timetable row", f"{a.time} {a.slot}: {a.med}", a.instruction, b.instruction))
    return rows


def add_sheet(wb, title, headers, rows, widths, verdict_cols=None, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).alignment = Alignment(wrap_text=True, vertical="top"); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers)); ws.row_dimensions[1].height = 62
        r0 = 2
    for c, h in enumerate(headers, 1):
        cell = ws.cell(r0, c, h); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor=TEAL); cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(rows, r0 + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(i, c, v); cell.alignment = Alignment(wrap_text=True, vertical="top")
        for c in (verdict_cols or []):
            ws.cell(i, c).fill = PatternFill("solid", fgColor=TINT)
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + c)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    last = r0 + len(rows)
    for col, lst in (verdict_cols or {}).items() if isinstance(verdict_cols, dict) else []:
        dv = DataValidation(type="list", formula1=lst, allow_blank=True); ws.add_data_validation(dv); dv.add(f"{chr(64 + col)}{r0 + 1}:{chr(64 + col)}{last}")
    return ws


def ai_rows(lang):
    en = {m.name: e for m, e in zip(RX.medications, explain(RX, "en"))}
    items = explain(RX, lang)
    out = []
    for m, e in zip(RX.medications, items):
        out.append((m.name, f"{en[m.name].purpose} / {en[m.name].watch_for}", "what it is for", e.purpose))
        out.append((m.name, "", "what to watch for", e.watch_for))
    return out


def main():
    wb = Workbook(); ws = wb.active; ws.title = "Start here"
    text = [("Dawa Buddy: Hindi and Kannada review", True),
            (f"Prepared {datetime.date.today()}. Thank you for helping. This takes about 45 minutes per language.", False), ("", False),
            ("What this app does", True),
            ("It reads a medicine bill and a prescription and tells a patient what to ask a seller or doctor. In Hindi and Kannada it explains each medicine, the safety notes and a daily timetable. "
             "Medicine names, ingredient names, numbers and units (mg, ml) stay as printed on the strip, in English letters, on purpose.", False), ("", False),
            ("What to do", True),
            ("1. Open the sheet for your language (Hindi: 'Hindi fixed text' and 'Hindi AI text'. Kannada: 'Kannada fixed text', 'Kannada medicine table', 'Kannada AI text').", False),
            ("2. Read the English meaning and the Hindi/Kannada text next to it. In the Verdict column choose OK, Change it, or Not sure.", False),
            ("3. If Change it, type the wording you would use in 'Your correction'. Short, plain, spoken-language wording is better than formal wording.", False), ("", False),
            ("Judge four things", True),
            ("Meaning: does it say the same as the English? Safe: it never tells the patient to stop or change a dose, never diagnoses. Plain: could an ordinary family member understand it? Natural: would a native speaker say it this way?", False), ("", False),
            ("The AI text sheets", True),
            ("A model wrote these, so they vary. Please also mark anything that is fluent but is NOT about the medicine (nonsense). For Kannada this is switched off in the app because a test produced nonsense; "
             "your verdict tells us whether it could ever be switched on.", False), ("", False),
            ("The 'Medicine map' sheet (for a pharmacist or doctor, in English)", True),
            ("It lists, per ingredient, what the app says it is for, the side effects it lists and the warnings it gives. Please check the medical meaning: is anything wrong, missing, or unsafe?", False)]
    for i, (t, bold) in enumerate(text, 1):
        c = ws.cell(i, 1, t); c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(bold=bold, size=13 if bold else 11, color=TEAL if bold else "000000")
    ws.column_dimensions["A"].width = 120

    for lang, name in (("hi", "Hindi"), ("kn", "Kannada")):
        rows = [(f"{name[0]}{i:03d}", part, where, en, txt, "", "", "") for i, (part, where, en, txt) in enumerate(fixed_rows(lang), 1)]
        ws = add_sheet(wb, f"{name} fixed text", ["ID", "Part", "Where it appears", "English meaning", name, "Verdict", "Your correction", "Notes"], rows,
                       [8, 22, 38, 52, 52, 14, 46, 30], verdict_cols={6: VERDICTS},
                       note=f"Fixed {name} text: never written by a model. Each row is one phrase or sentence the app shows. Choose a Verdict for every row; add your wording under 'Your correction' when it should change.")
    medrows = []
    for k, label in (("PURPOSE", "Sentence frame: what it is for"), ("SIDE_SENTENCE", "Sentence frame: side effects"), ("PURPOSE_JOIN", "Joiner between two purposes")):
        medrows.append((label, k, getattr(medinfo, k)["en"], getattr(medinfo, k)["kn"]))
    medrows += [("What it is for (completes 'This medicine is used ...')", k, en, kn) for k, (en, kn) in medinfo._P.items()]
    medrows += [("Side effect (completes 'Some people get ...')", k, en, kn) for k, (en, kn) in medinfo.SIDE.items()]
    medrows += [("Warning (whole sentence)", k, en, kn) for k, (en, kn) in medinfo.WARN.items()]
    add_sheet(wb, "Kannada medicine table", ["ID", "Kind", "Piece", "English meaning", "Kannada", "Verdict", "Your correction", "Notes"],
              [(f"KM{i:03d}", a, b, c, d, "", "", "") for i, (a, b, c, d) in enumerate(medrows, 1)], [8, 34, 20, 52, 52, 14, 46, 30], verdict_cols={6: VERDICTS},
              note="The fixed table that replaces AI text for common ingredients. Every Kannada medicine explanation is built from these short pieces, so checking these checks all of them.")
    mapping = [(name, medinfo._P[pk][0], ", ".join(medinfo.SIDE[x][0] for x in se) or "-", " / ".join(medinfo.WARN[x][0] for x in wn) or "-", "", "")
               for name, (pk, se, wn) in medinfo.MEDS.items()]
    add_sheet(wb, "Medicine map", ["Ingredient", "What the app says it is for", "Side effects it lists", "Warnings it gives", "Medically OK?", "Correction or missing warning"], mapping,
              [26, 52, 52, 70, 16, 46], verdict_cols={5: VERDICTS},
              note=f"For a pharmacist or doctor (English). {len(mapping)} ingredients. Is the purpose right? Are the common side effects right? Is a warning wrong, missing or unsafe? Combinations are built from these rows.")
    if llm.is_offline():
        print("NEBIUS_API_KEY not set: the AI-text sheets are skipped")
    else:
        import os
        for lang, name, note in (("hi", "Hindi", "Written by a model (Nemotron Ultra) and shown to Hindi readers, labelled as AI. Please rate each line."),
                                 ("kn", "Kannada", "Written by a model and NOT shown to Kannada readers (a test produced nonsense). Rate each line so we know whether it could ever be switched on.")):
            if lang == "kn":
                os.environ["DAWA_KN_MODEL_TEXT"] = "1"
                _info = medinfo.info; medinfo.info = lambda *a, **k: None       # show the raw model output, not the fixed table that replaces it
            rows = [(med, en, kind, txt, "", "", "", "", "", "") for med, en, kind, txt in ai_rows(lang)]
            if lang == "kn":
                medinfo.info = _info
            add_sheet(wb, f"{name} AI text" + (" (off)" if lang == "kn" else ""), ["Medicine", "English (model's English version)", "Kind", name, "Meaning right?", "Safe?", "Plain?", "Natural?", "Nonsense?", "Notes"],
                      rows, [16, 46, 18, 56, 14, 10, 10, 10, 12, 30], verdict_cols={5: YN, 6: YN, 7: YN, 8: YN, 9: YN}, note=note)
    out = pathlib.Path("my_data") / "hindi_kannada_review.xlsx"
    wb.save(out)
    print("written", out, "|", ", ".join(wb.sheetnames))


main()
